"""Create verified, readable Turkish document artifacts without changing the source."""

from __future__ import annotations

import html
import os
import re
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

from ksi_local.atomic_files import atomic_write_bytes, atomic_write_text


@dataclass(frozen=True)
class DocumentOutputResult:
    markdown_path: Path
    docx_path: Path
    pdf_path: Path
    text_path: Path
    block_count: int
    embedded_pdf_font: bool


def _text(record: dict[str, Any]) -> str:
    value = record.get("translated_text")
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Türkçe belge çıktısında boş veya geçersiz blok var.")
    return value


def _table_key(record: dict[str, Any]) -> tuple[object, ...] | None:
    if record.get("block_type") != "table_cell":
        return None
    location = record.get("location")
    if not isinstance(location, dict):
        raise ValueError("Türkçe belge tablo konumu geçersiz.")
    return (
        location.get("part", "body"),
        location.get("section"),
        location.get("parent_table"),
        location.get("table"),
        location.get("depth", 1),
    )


def _groups(records: list[dict[str, Any]]) -> list[list[dict[str, Any]]]:
    result: list[list[dict[str, Any]]] = []
    index = 0
    while index < len(records):
        key = _table_key(records[index])
        if key is None:
            result.append([records[index]])
            index += 1
            continue
        group: list[dict[str, Any]] = []
        while index < len(records) and _table_key(records[index]) == key:
            group.append(records[index])
            index += 1
        result.append(group)
    return result


def _table_matrix(group: list[dict[str, Any]]) -> tuple[list[list[str]], bool]:
    positions: dict[tuple[int, int], str] = {}
    max_row = 0
    max_column = 0
    has_header = False
    for record in group:
        location = record["location"]
        row = location.get("row")
        column = location.get("column")
        if not isinstance(row, int) or row < 1 or not isinstance(column, int) or column < 1:
            raise ValueError("Türkçe belge tablo satır/sütun konumu geçersiz.")
        if (row, column) in positions:
            raise ValueError("Türkçe belge tablosunda yinelenmiş hücre var.")
        positions[(row, column)] = _text(record)
        max_row = max(max_row, row)
        max_column = max(max_column, column)
        style = record.get("style")
        has_header = has_header or bool(isinstance(style, dict) and style.get("header"))
    matrix = [
        [positions.get((row, column), "") for column in range(1, max_column + 1)]
        for row in range(1, max_row + 1)
    ]
    return matrix, has_header


def _markdown_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def render_markdown(records: list[dict[str, Any]], *, source_title: str) -> str:
    lines = ["# Türkçe Belge", "", f"Kaynak: {source_title.strip()[:500] or 'Yerel belge'}", ""]
    for group in _groups(records):
        first = group[0]
        if _table_key(first) is not None:
            matrix, has_header = _table_matrix(group)
            header = matrix[0]
            lines.append("| " + " | ".join(_markdown_cell(item) for item in header) + " |")
            lines.append("| " + " | ".join("---" for _ in header) + " |")
            body = matrix[1:] if has_header or len(matrix) > 1 else []
            for row in body:
                lines.append("| " + " | ".join(_markdown_cell(item) for item in row) + " |")
            lines.append("")
            continue
        record = first
        value = _text(record)
        block_type = str(record.get("block_type") or "paragraph")
        style = record.get("style") if isinstance(record.get("style"), dict) else {}
        if block_type == "heading":
            level = min(max(int(style.get("heading_level") or 1), 1), 6)
            lines.extend((f"{'#' * level} {value}", ""))
        elif block_type == "list_item":
            indent = max(int(style.get("indent") or 0) // 2, 0)
            lines.extend((f"{'  ' * indent}- {value}", ""))
        elif block_type == "code_block":
            language = re.sub(r"[^A-Za-z0-9_+.-]", "", str(style.get("language") or ""))[:30]
            fence = "````" if "```" in value else "```"
            lines.extend((f"{fence}{language}", value, fence, ""))
        elif block_type == "link_definition" and style.get("target"):
            lines.extend((f"[{value}]: {style['target']}", ""))
        elif block_type == "footnote":
            footnote_id = record.get("location", {}).get("footnote_id", "")
            lines.extend((f"[^dipnot-{footnote_id}]: {value}", ""))
        else:
            lines.extend((value, ""))
    return "\n".join(lines).rstrip() + "\n"


def render_text(records: list[dict[str, Any]], *, source_title: str) -> str:
    lines = ["KSI Local Studio Türkçe Belge", f"Kaynak: {source_title.strip()[:500] or 'Yerel belge'}", ""]
    lines.extend(_text(record) + "\n" for record in records)
    return "\n".join(lines).rstrip() + "\n"


def _docx_bytes(records: list[dict[str, Any]], *, source_title: str) -> bytes:
    document = Document()
    document.core_properties.title = "Türkçe Belge"
    document.core_properties.subject = source_title.strip()[:500]
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    document.add_heading("Türkçe Belge", level=0)
    document.add_paragraph(f"Kaynak: {source_title.strip()[:500] or 'Yerel belge'}")
    for group in _groups(records):
        first = group[0]
        if _table_key(first) is not None:
            matrix, has_header = _table_matrix(group)
            table = document.add_table(rows=len(matrix), cols=len(matrix[0]))
            table.style = "Table Grid"
            for row_index, row in enumerate(matrix):
                for column_index, value in enumerate(row):
                    cell = table.cell(row_index, column_index)
                    cell.text = value
                    if has_header and row_index == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
            continue
        record = first
        value = _text(record)
        block_type = str(record.get("block_type") or "paragraph")
        style = record.get("style") if isinstance(record.get("style"), dict) else {}
        if block_type == "heading":
            document.add_heading(value, level=min(max(int(style.get("heading_level") or 1), 1), 9))
        elif block_type == "list_item":
            document.add_paragraph(value, style="List Bullet")
        elif block_type == "code_block":
            paragraph = document.add_paragraph()
            run = paragraph.add_run(value)
            run.font.name = "Menlo"
            run.font.size = Pt(9)
        elif block_type == "footnote":
            paragraph = document.add_paragraph()
            paragraph.add_run("Dipnot: ").bold = True
            paragraph.add_run(value)
        else:
            document.add_paragraph(value)
    buffer = BytesIO()
    document.save(buffer)
    payload = buffer.getvalue()
    Document(BytesIO(payload))
    return payload


def _html_document(records: list[dict[str, Any]], *, source_title: str) -> str:
    body = ["<h1>Türkçe Belge</h1>", f"<p><b>Kaynak:</b> {html.escape(source_title)}</p>"]
    for group in _groups(records):
        first = group[0]
        if _table_key(first) is not None:
            matrix, has_header = _table_matrix(group)
            body.append("<table>")
            for row_index, row in enumerate(matrix):
                tag = "th" if has_header and row_index == 0 else "td"
                cells = "".join(
                    f"<{tag}>{html.escape(value)}</{tag}>" for value in row
                )
                body.append(f"<tr>{cells}</tr>")
            body.append("</table>")
            continue
        value = html.escape(_text(first)).replace("\n", "<br>")
        block_type = str(first.get("block_type") or "paragraph")
        style = first.get("style") if isinstance(first.get("style"), dict) else {}
        if block_type == "heading":
            level = min(max(int(style.get("heading_level") or 1), 1), 6)
            body.append(f"<h{level}>{value}</h{level}>")
        elif block_type == "list_item":
            body.append(f"<ul><li>{value}</li></ul>")
        elif block_type == "code_block":
            body.append(f"<pre>{value}</pre>")
        else:
            body.append(f"<p>{value}</p>")
    css = """
    body { font-family: Arial, sans-serif; font-size: 10.5pt; line-height: 1.35; }
    h1 { font-size: 22pt; } h2 { font-size: 17pt; } h3 { font-size: 14pt; }
    table { border-collapse: collapse; width: 100%; margin: 8pt 0; }
    th, td { border: 1px solid #777; padding: 4pt; vertical-align: top; }
    th { background: #e9eef2; } pre { font-family: Menlo, monospace; font-size: 8.5pt; }
    """
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<style>{css}</style></head><body>{''.join(body)}</body></html>"
    )


def _pdf_has_embedded_font(reader: PdfReader) -> bool:
    def descriptor_embedded(font: object) -> bool:
        resolved = font.get_object()
        descriptor = resolved.get("/FontDescriptor")
        if descriptor is not None:
            value = descriptor.get_object()
            if any(value.get(key) is not None for key in ("/FontFile", "/FontFile2", "/FontFile3")):
                return True
        descendants = resolved.get("/DescendantFonts") or []
        return any(descriptor_embedded(item) for item in descendants)

    for page in reader.pages:
        resources = page.get("/Resources")
        if resources is None:
            continue
        fonts = resources.get_object().get("/Font")
        if fonts is not None and any(
            descriptor_embedded(item) for item in fonts.get_object().values()
        ):
            return True
    return False


def _write_pdf(path: Path, records: list[dict[str, Any]], *, source_title: str) -> bool:
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QFont, QGuiApplication, QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    app = QGuiApplication.instance()
    created = app is None
    if created:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        app = QGuiApplication(["ksi_local-document-output"])
    try:
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(path))
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        printer.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Unit.Millimeter)
        document = QTextDocument()
        document.setDefaultFont(QFont("Arial", 11))
        document.setHtml(_html_document(records, source_title=source_title))
        document.print_(printer)
        reader = PdfReader(path, strict=True)
        if not reader.pages:
            raise RuntimeError("Türkçe belge PDF çıktısı boş.")
        embedded = _pdf_has_embedded_font(reader)
        if not embedded:
            raise RuntimeError("Türkçe belge PDF fontu gömülü değil.")
        path.chmod(0o600)
        return embedded
    finally:
        if created:
            app.quit()


def create_document_outputs(
    records: list[dict[str, Any]], output_directory: str | Path, *, source_title: str
) -> DocumentOutputResult:
    if not records:
        raise ValueError("Türkçe belge çıktısı için blok bulunamadı.")
    output_candidate = Path(output_directory).expanduser()
    if output_candidate.is_symlink():
        raise ValueError("Türkçe belge çıktı klasörü sembolik bağlantı olamaz.")
    output = output_candidate.resolve()
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    staging = Path(tempfile.mkdtemp(prefix=".belge-turkce-", dir=output))
    targets = {
        "markdown": output / "belge-turkce.md",
        "docx": output / "belge-turkce.docx",
        "pdf": output / "belge-turkce.pdf",
        "text": output / "belge-turkce.txt",
    }
    try:
        markdown = render_markdown(records, source_title=source_title)
        text = render_text(records, source_title=source_title)
        atomic_write_text(staging / targets["markdown"].name, markdown)
        atomic_write_text(staging / targets["text"].name, text)
        atomic_write_bytes(
            staging / targets["docx"].name,
            _docx_bytes(records, source_title=source_title),
        )
        embedded = _write_pdf(staging / targets["pdf"].name, records, source_title=source_title)
        if not (staging / targets["markdown"].name).read_text(encoding="utf-8").strip():
            raise RuntimeError("Türkçe Markdown çıktısı yeniden okunamadı.")
        Document(staging / targets["docx"].name)
        PdfReader(staging / targets["pdf"].name, strict=True)
        for name, target in targets.items():
            os.replace(staging / target.name, target)
    finally:
        import shutil

        shutil.rmtree(staging, ignore_errors=True)
    return DocumentOutputResult(
        markdown_path=targets["markdown"],
        docx_path=targets["docx"],
        pdf_path=targets["pdf"],
        text_path=targets["text"],
        block_count=len(records),
        embedded_pdf_font=embedded,
    )
