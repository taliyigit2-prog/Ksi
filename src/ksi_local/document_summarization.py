"""Traceable, resumable Turkish summaries for canonical document blocks."""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from collections.abc import Callable
from dataclasses import asdict, dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Pt
from pypdf import PdfReader

from ksi_local.atomic_files import atomic_write_bytes, atomic_write_json, atomic_write_text
from ksi_local.document_translation import _canonical_blocks
from ksi_local.ollama_client import OllamaClient


SCHEMA_VERSION = 1
SUMMARY_POLICY_VERSION = 5
SUMMARY_MODEL = "qwen3.5:4b"
MAX_TRANSLATION_BYTES = 100 * 1024 * 1024
SUMMARY_SOURCE_CHOICES = {"auto", "source", "translation"}
PROFILE_SETTINGS = {
    "short": {
        "label": "Kısa",
        "claims_per_chunk": 3,
        "short_summary": 2,
        "main_topics": 3,
        "important_ideas": 3,
        "conclusions": 2,
        "actions": 2,
    },
    "standard": {
        "label": "Standart",
        "claims_per_chunk": 6,
        "short_summary": 3,
        "main_topics": 6,
        "important_ideas": 6,
        "conclusions": 4,
        "actions": 4,
    },
    "detailed": {
        "label": "Ayrıntılı",
        "claims_per_chunk": 8,
        "short_summary": 5,
        "main_topics": 10,
        "important_ideas": 12,
        "conclusions": 8,
        "actions": 8,
    },
}
CATEGORIES = {"ana_fikir", "önemli_nokta", "sonuç", "eylem"}
SECTIONS = ("short_summary", "main_idea", "important_points", "conclusions", "actions")


@dataclass(frozen=True)
class DocumentEvidence:
    block_id: str
    quote: str

    def to_dict(self) -> dict[str, str]:
        return asdict(self)


@dataclass(frozen=True)
class DocumentClaim:
    id: str
    text: str
    category: str
    block_ids: tuple[str, ...]
    evidence: tuple[DocumentEvidence, ...]

    def to_dict(self) -> dict[str, object]:
        return {
            "id": self.id,
            "text": self.text,
            "category": self.category,
            "block_ids": list(self.block_ids),
            "evidence": [item.to_dict() for item in self.evidence],
        }


@dataclass(frozen=True)
class DocumentSummaryProgress:
    completed: int
    total: int


@dataclass(frozen=True)
class DocumentSummaryResult:
    markdown_path: Path
    docx_path: Path
    pdf_path: Path | None
    trace_path: Path
    quality_path: Path
    checkpoint_path: Path
    profile: str
    source_mode: str
    chunk_count: int
    claim_count: int
    statement_count: int
    quality: dict[str, object]


def _ground_claim_text(
    text: str, category: str, evidence: list[DocumentEvidence] | tuple[DocumentEvidence, ...]
) -> tuple[str, str]:
    evidence_text = " ".join(item.quote for item in evidence)
    unsupported_event_stems = (
        "oluşturul", "hazırlan", "belgelendir", "alınd", "yapıl", "uygulan",
        "gerçekleş", "planlan", "belirlen", "sınırlandır", "sunul", "kaydedil",
    )
    negation = re.compile(
        r"\b(?:değil\w*|yok\w*|hiçbir|yasak\w*|izin vermez|yapmamalı\w*|açmayın|\w+(?:ma|me)d[ıiuü])\b",
        re.I,
    )
    numbers = lambda value: set(re.findall(r"\d+(?:[.,]\d+)*%?", value))
    unsupported_event = any(stem in text.casefold() for stem in unsupported_event_stems) and not any(
        stem in evidence_text.casefold() for stem in unsupported_event_stems
    )
    lost_critical_fact = bool(negation.search(evidence_text)) and (
        not negation.search(text) or not numbers(evidence_text) <= numbers(text)
    )
    if unsupported_event or lost_critical_fact:
        text = "; ".join(dict.fromkeys(item.quote for item in evidence))
        if len(text) > 800:
            text = text[:797].rstrip() + "..."
        category = "önemli_nokta"
    return text, category


def _sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def _sha256_text(value: str) -> str:
    return _sha256_bytes(value.encode("utf-8"))


def _normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip().casefold()


def _evidence_tokens(value: str) -> list[str]:
    return re.findall(r"[^\W_]+(?:[.,][0-9]+)?", value.casefold(), flags=re.UNICODE)


def _exact_or_repaired_evidence_quote(quote: str, source: str) -> str | None:
    """Return source-exact evidence when a small model only reformats its quote.

    Evidence stays verbatim: the repair may select the complete source block, but only
    when the proposed quote has strong lexical overlap, introduces no number, and does
    not change whether the passage contains a negation. Unrelated/hallucinated quotes
    remain invalid.
    """

    if _normalized(quote) in _normalized(source):
        return quote
    if len(source) > 500:
        return None
    quote_tokens = _evidence_tokens(quote)
    source_tokens = _evidence_tokens(source)
    if len(quote_tokens) < 4:
        return None
    source_token_set = set(source_tokens)
    overlap = sum(token in source_token_set for token in quote_tokens) / len(quote_tokens)
    quote_numbers = set(re.findall(r"\d+(?:[.,]\d+)*", quote))
    source_numbers = set(re.findall(r"\d+(?:[.,]\d+)*", source))
    negation = re.compile(r"\b(?:değil|yok|hayır|not|no|never|without)\b", re.I)
    if (
        overlap >= 0.8
        and quote_numbers <= source_numbers
        and bool(negation.search(quote)) == bool(negation.search(source))
    ):
        return source
    return None


def _translation_blocks(
    path: str | Path,
    quality_path: str | Path,
    canonical: list[dict[str, Any]],
) -> tuple[str, list[dict[str, Any]]]:
    candidate = Path(path).expanduser()
    quality_candidate = Path(quality_path).expanduser()
    if candidate.is_symlink() or quality_candidate.is_symlink():
        raise ValueError("Belge çevirisi veya kalite raporu sembolik bağlantı olamaz.")
    resolved = candidate.resolve(strict=True)
    resolved_quality = quality_candidate.resolve(strict=True)
    if not resolved.is_file() or resolved.stat().st_size > MAX_TRANSLATION_BYTES:
        raise ValueError("Doğrulanmış belge çevirisi bulunamadı veya çok büyük.")
    try:
        quality = json.loads(resolved_quality.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError("Belge çeviri kalite raporu okunamadı.") from error
    if not isinstance(quality, dict) or quality.get("passed") is not True:
        raise ValueError("Belge özeti için çeviri kalite raporunun geçmiş olması gerekir.")
    raw = resolved.read_bytes()
    records: list[dict[str, Any]] = []
    for line in raw.decode("utf-8").splitlines():
        try:
            record = json.loads(line)
        except json.JSONDecodeError as error:
            raise ValueError("Belge çeviri JSONL yapısı geçersiz.") from error
        if not isinstance(record, dict):
            raise ValueError("Belge çeviri satırı nesne biçiminde değil.")
        records.append(record)
    if len(records) != len(canonical):
        raise ValueError("Belge çevirisi kaynakla aynı blok sayısını taşımıyor.")
    for source, translated in zip(canonical, records, strict=True):
        source_text = str(source["text"])
        target_text = translated.get("translated_text")
        if (
            translated.get("id") != source["id"]
            or translated.get("sequence") != source["sequence"]
            or translated.get("source_text") != source_text
            or translated.get("source_block_sha256") != _sha256_text(source_text)
            or not isinstance(target_text, str)
            or not target_text.strip()
            or translated.get("translation_sha256") != _sha256_text(target_text)
        ):
            raise ValueError("Belge çeviri blokları kaynakla veya SHA-256 ile eşleşmiyor.")
    return _sha256_bytes(raw), records


def _summary_blocks(
    canonical_path: str | Path,
    *,
    source_mode: str,
    translation_path: str | Path | None,
    translation_quality_path: str | Path | None,
) -> tuple[str, str, str | None, list[dict[str, Any]]]:
    if source_mode not in SUMMARY_SOURCE_CHOICES:
        raise ValueError("Belge özet kaynağı otomatik, kaynak veya çeviri olmalıdır.")
    _canonical, canonical_sha256, blocks = _canonical_blocks(canonical_path)
    translated_available = bool(translation_path and translation_quality_path)
    resolved_mode = (
        "translation"
        if source_mode == "translation" or (source_mode == "auto" and translated_available)
        else "source"
    )
    translation_sha256: str | None = None
    if resolved_mode == "translation":
        if not translation_path or not translation_quality_path:
            raise ValueError("Türkçe çeviriden özet için doğrulanmış çeviri bulunamadı.")
        translation_sha256, translated = _translation_blocks(
            translation_path, translation_quality_path, blocks
        )
        chosen = [
            {
                "id": source["id"],
                "sequence": source["sequence"],
                "block_type": source["block_type"],
                "location": source["location"],
                "text": target["translated_text"],
                "original_text": source["text"],
            }
            for source, target in zip(blocks, translated, strict=True)
        ]
    else:
        chosen = [
            {
                "id": block["id"],
                "sequence": block["sequence"],
                "block_type": block["block_type"],
                "location": block["location"],
                "text": block["text"],
                "original_text": block["text"],
            }
            for block in blocks
        ]
    return resolved_mode, canonical_sha256, translation_sha256, chosen


def split_document_blocks(
    blocks: list[dict[str, Any]], *, max_chars: int = 8000, max_blocks: int = 8
) -> list[list[dict[str, Any]]]:
    if max_chars < 1000 or max_chars > 20_000:
        raise ValueError("Belge özet parçası 1000–20000 karakter arasında olmalıdır.")
    if max_blocks < 1 or max_blocks > 100:
        raise ValueError("Belge özet parçası 1–100 blok arasında olmalıdır.")
    chunks: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    length = 0
    for block in blocks:
        block_length = len(str(block["text"])) + 120
        if current and (length + block_length > max_chars or len(current) >= max_blocks):
            chunks.append(current)
            current = []
            length = 0
        current.append(block)
        length += block_length
    if current:
        chunks.append(current)
    return chunks


def _claim_schema(block_ids: list[str], limit: int, minimum: int) -> dict[str, object]:
    return {
        "type": "object",
        "properties": {
            "claims": {
                "type": "array",
                "minItems": minimum,
                "maxItems": limit,
                "items": {
                    "type": "object",
                    "properties": {
                        "text": {"type": "string", "maxLength": 300},
                        "category": {"type": "string", "enum": sorted(CATEGORIES)},
                        "block_ids": {
                            "type": "array",
                            "minItems": 1,
                            "maxItems": 12,
                            "uniqueItems": True,
                            "items": {"type": "string", "enum": block_ids},
                        },
                        "evidence": {
                            "type": "array",
                            "minItems": 1,
                            "maxItems": 12,
                            "items": {
                                "type": "object",
                                "properties": {
                                    "block_id": {"type": "string", "enum": block_ids},
                                    "quote": {"type": "string", "maxLength": 280},
                                },
                                "required": ["block_id", "quote"],
                                "additionalProperties": False,
                            },
                        },
                    },
                    "required": ["text", "category", "block_ids", "evidence"],
                    "additionalProperties": False,
                },
            }
        },
        "required": ["claims"],
        "additionalProperties": False,
    }


def _payload(raw: str) -> dict[str, Any]:
    cleaned = raw.strip()
    if cleaned.startswith("```"):
        cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.I)
        cleaned = re.sub(r"\s*```$", "", cleaned)
    try:
        value = json.loads(cleaned)
    except json.JSONDecodeError as error:
        raise RuntimeError("Belge özet modeli geçerli JSON döndürmedi.") from error
    if not isinstance(value, dict):
        raise RuntimeError("Belge özet modeli JSON nesnesi döndürmedi.")
    return value


def _parse_claims(
    raw: str,
    *,
    blocks: list[dict[str, Any]],
    first_number: int,
    limit: int,
    minimum: int,
) -> list[DocumentClaim]:
    records = _payload(raw).get("claims")
    if not isinstance(records, list) or not minimum <= len(records) <= limit:
        raise RuntimeError("Belge özet modeli geçerli sayıda kanıt iddiası döndürmedi.")
    block_map = {str(item["id"]): str(item["text"]) for item in blocks}
    claims: list[DocumentClaim] = []
    seen: set[str] = set()
    for record in records:
        if not isinstance(record, dict):
            raise RuntimeError("Belge özet iddiası nesne biçiminde değil.")
        text = re.sub(r"\s+", " ", str(record.get("text") or "")).strip()
        category = str(record.get("category") or "")
        raw_ids = record.get("block_ids")
        raw_evidence = record.get("evidence")
        if (
            not text
            or len(text) > 800
            or category not in CATEGORIES
            or not isinstance(raw_ids, list)
            or not raw_ids
            or not isinstance(raw_evidence, list)
            or not raw_evidence
        ):
            raise RuntimeError("Belge özet iddiasının zorunlu alanları geçersiz.")
        block_ids = tuple(dict.fromkeys(str(item) for item in raw_ids))
        if len(block_ids) != len(raw_ids) or any(item not in block_map for item in block_ids):
            raise RuntimeError("Belge özet iddiasında bilinmeyen veya yinelenmiş blok kimliği var.")
        evidence: list[DocumentEvidence] = []
        evidence_ids: set[str] = set()
        for item in raw_evidence:
            if not isinstance(item, dict):
                raise RuntimeError("Belge özet kanıtı nesne biçiminde değil.")
            block_id = str(item.get("block_id") or "")
            quote = re.sub(r"\s+", " ", str(item.get("quote") or "")).strip()
            repaired_quote = (
                _exact_or_repaired_evidence_quote(quote, block_map[block_id])
                if block_id in block_map and quote and len(quote) <= 500
                else None
            )
            if block_id not in block_ids or repaired_quote is None:
                raise RuntimeError(
                    "Belge özet kanıtı kaynak blokta birebir bulunamadı veya kimliği geçersiz."
                )
            evidence.append(DocumentEvidence(block_id, repaired_quote))
            evidence_ids.add(block_id)
        if evidence_ids != set(block_ids):
            raise RuntimeError("Belge özet iddiasının her kaynak bloğu için kanıt alıntısı gerekli.")
        text, category = _ground_claim_text(text, category, evidence)
        key = text.casefold()
        if key in seen:
            raise RuntimeError("Belge özet modeli yinelenmiş iddia döndürdü.")
        seen.add(key)
        claims.append(
            DocumentClaim(
                id=f"C{first_number + len(claims):06d}",
                text=text,
                category=category,
                block_ids=block_ids,
                evidence=tuple(evidence),
            )
        )
    return claims


def _generate_claims(
    chunk: list[dict[str, Any]],
    *,
    client: OllamaClient,
    model: str,
    profile: str,
    chunk_index: int,
    chunk_count: int,
    first_number: int,
) -> list[DocumentClaim]:
    limit = int(PROFILE_SETTINGS[profile]["claims_per_chunk"])
    desired_minimum = {"short": 2, "standard": 6, "detailed": 8}[profile]
    minimum = min(limit, len(chunk), desired_minimum)
    schema = _claim_schema([str(item["id"]) for item in chunk], limit, minimum)
    input_blocks = [
        {
            "id": item["id"],
            "type": item["block_type"],
            "location": item["location"],
            "text": item["text"],
        }
        for item in chunk
    ]
    prompt = (
        f"Bu belgenin {chunk_index + 1}/{chunk_count} parçasıdır. Yalnız INPUT içeriğine dayanarak "
        f"en az {minimum}, en fazla {limit} Türkçe kanıt iddiası üret. Dış bilgi, tahmin veya yeni sayı ekleme. "
        "Tarih, kişi, saat, sayı, sınır, yasak, iletişim ve yapılacak iş gibi birbirinden farklı "
        "somut bilgileri tek bir genel cümlede birleştirme; ayrı iddialar halinde yaz. "
        "Kaynak yalnız bir tarih etiketi veya belge kopyası tarihi veriyorsa o tarihte bir olayın "
        "gerçekleştiğini, belgenin oluşturulduğunu ya da işlemin yapıldığını varsayma. "
        "Her iddiayı doğrudan destekleyen bütün block_ids değerlerini yaz. Her block_id için, "
        "o bloğun text alanında boşluklar dışında birebir bulunan kısa bir evidence quote kopyala. "
        "Alıntıyı çevirme veya düzeltme. Açık bir sonuç ya da yapılacak iş yoksa o kategoride iddia "
        "üretme. JSON şeması dışına çıkma.\nZORUNLU JSON ŞEMASI:\n"
        + json.dumps(schema, ensure_ascii=False)
        + "\nINPUT:\n"
        + json.dumps(input_blocks, ensure_ascii=False)
    )
    last_error: RuntimeError | None = None
    for attempt in range(3):
        raw = client.generate(
            model=model,
            prompt=prompt,
            system=(
                "Sen kaynak sadakatine öncelik veren Türkçe belge editörüsün. Yalnız doğrulanabilir "
                "iddialar yaz ve geçerli JSON dışında hiçbir metin döndürme."
            ),
            json_schema=schema,
            temperature=0.0 if attempt == 0 else 0.1 * attempt,
            keep_alive="5m",
            max_tokens=1600,
        )
        try:
            return _parse_claims(
                raw, blocks=chunk, first_number=first_number, limit=limit, minimum=minimum
            )
        except RuntimeError as error:
            last_error = error
    raise last_error or RuntimeError("Belge özet modeli doğrulanabilir iddia üretmedi.")


def _final_schema(claim_ids: list[str], profile: str) -> dict[str, object]:
    settings = PROFILE_SETTINGS[profile]
    properties: dict[str, object] = {}
    for section in SECTIONS:
        limit_key = {
            "main_idea": "main_topics",
            "important_points": "important_ideas",
        }.get(section, section)
        field: dict[str, object] = {
            "type": "array",
            "maxItems": int(settings[limit_key]),
            "uniqueItems": True,
            "items": {"type": "string", "enum": claim_ids},
        }
        if section in {"short_summary", "main_idea"}:
            field["minItems"] = 1
        properties[section] = field
    return {
        "type": "object",
        "properties": properties,
        "required": list(SECTIONS),
        "additionalProperties": False,
    }


def _fallback_sections(
    claims: list[DocumentClaim], profile: str
) -> dict[str, list[str]]:
    settings = PROFILE_SETTINGS[profile]

    def pick(category: str, section: str) -> list[str]:
        return [
            item.id for item in claims if item.category == category
        ][: int(settings[section])]

    main = pick("ana_fikir", "main_topics") or [
        item.id for item in claims[: int(settings["main_topics"])]
    ]
    return {
        "short_summary": [item.id for item in claims[: int(settings["short_summary"])]],
        "main_idea": main,
        "important_points": pick("önemli_nokta", "important_ideas"),
        "conclusions": pick("sonuç", "conclusions"),
        "actions": pick("eylem", "actions"),
    }


def _refine_sections(
    sections: dict[str, list[str]], claims: list[DocumentClaim], profile: str
) -> dict[str, list[str]]:
    """Keep model selection useful while deterministically retaining critical facts."""

    claim_map = {item.id: item for item in claims}

    def informative(claim: DocumentClaim) -> bool:
        text = claim.text.strip()
        return bool(re.search(r"\d|@", text) or len(re.findall(r"\w+", text, re.UNICODE)) >= 5)

    def priority(claim: DocumentClaim) -> tuple[int, int]:
        text = (claim.text + " " + " ".join(item.quote for item in claim.evidence)).casefold()
        score = 0
        if re.search(
            r"\b(?:değil\w*|yok\w*|hiçbir|yasak\w*|izin vermez|yapmamalı\w*|açmayın|\w+(?:ma|me)d[ıiuü])\b",
            text,
        ):
            score += 12
        if "@" in text:
            score += 10
        if re.search(r"\d", text):
            score += 6
        if claim.category == "eylem":
            score += 4
        if claim.category == "önemli_nokta":
            score += 3
        return score, -int(claim.id[1:])

    candidates = sorted((item for item in claims if informative(item)), key=priority, reverse=True)
    refined: dict[str, list[str]] = {section: [] for section in SECTIONS}
    expected_category = {
        "important_points": "önemli_nokta",
        "conclusions": "sonuç",
        "actions": "eylem",
    }
    used: set[str] = set()
    for section in SECTIONS:
        limit_key = {"main_idea": "main_topics", "important_points": "important_ideas"}.get(
            section, section
        )
        limit = int(PROFILE_SETTINGS[profile][limit_key])
        values: list[str] = []
        source_sets: set[frozenset[str]] = set()
        for claim_id in sections.get(section, []):
            claim = claim_map.get(claim_id)
            if claim is None or not informative(claim):
                continue
            required = expected_category.get(section)
            if required and claim.category != required:
                continue
            if section != "short_summary" and claim_id in used:
                continue
            claim_sources = frozenset(claim.block_ids)
            if claim_sources in source_sets:
                continue
            if claim_id not in values:
                values.append(claim_id)
                source_sets.add(claim_sources)
                if section != "short_summary":
                    used.add(claim_id)
            if len(values) >= limit:
                break
        refined[section] = values

    refined["short_summary"] = []
    short_sources: set[frozenset[str]] = set()
    for claim in candidates:
        if len(refined["short_summary"]) >= int(PROFILE_SETTINGS[profile]["short_summary"]):
            break
        claim_sources = frozenset(claim.block_ids)
        if claim_sources in short_sources:
            continue
        refined["short_summary"].append(claim.id)
        short_sources.add(claim_sources)

    for section in ("main_idea",):
        limit_key = "short_summary" if section == "short_summary" else "main_topics"
        limit = int(PROFILE_SETTINGS[profile][limit_key])
        section_sources = {
            frozenset(claim_map[claim_id].block_ids)
            for claim_id in refined[section]
        }
        for claim in candidates:
            if len(refined[section]) >= limit:
                break
            if claim.id in refined[section] or (section != "short_summary" and claim.id in used):
                continue
            claim_sources = frozenset(claim.block_ids)
            if claim_sources in section_sources:
                continue
            refined[section].append(claim.id)
            section_sources.add(claim_sources)
            if section != "short_summary":
                used.add(claim.id)

    important_limit = int(PROFILE_SETTINGS[profile]["important_ideas"])
    for claim in candidates:
        if len(refined["important_points"]) >= important_limit:
            break
        if priority(claim)[0] < 6 or claim.id in used:
            continue
        refined["important_points"].append(claim.id)
        used.add(claim.id)
    return refined


def _select_sections(
    claims: list[DocumentClaim],
    *,
    client: OllamaClient,
    model: str,
    profile: str,
    unload_model_at_end: bool,
) -> dict[str, list[str]]:
    schema = _final_schema([item.id for item in claims], profile)
    prompt = (
        "Aşağıdaki doğrulanmış iddialardan Türkçe belge özeti için seçim yap. Yeni metin veya kimlik "
        "üretme; yalnız claim id kullan. short_summary ana iddiaları tekrar edebilir, diğer ayrıntılı "
        "bölümlerde aynı id yalnız bir kez yer alsın. Kaynakta sonuç veya eylem yoksa ilgili dizi boş "
        "kalsın.\nZORUNLU JSON ŞEMASI:\n"
        + json.dumps(schema, ensure_ascii=False)
        + "\nİDDİALAR:\n"
        + json.dumps([item.to_dict() for item in claims], ensure_ascii=False)
    )
    try:
        raw = client.generate(
            model=model,
            prompt=prompt,
            json_schema=schema,
            temperature=0.0,
            keep_alive=0 if unload_model_at_end else "5m",
            max_tokens=800,
        )
        payload = _payload(raw)
        allowed = {item.id for item in claims}
        result: dict[str, list[str]] = {}
        used: set[str] = set()
        for section in SECTIONS:
            records = payload.get(section)
            if not isinstance(records, list):
                raise RuntimeError("Belge özet final bölümü liste değil.")
            values = [str(item) for item in records]
            if len(values) != len(set(values)) or any(item not in allowed for item in values):
                raise RuntimeError("Belge özet finalinde bilinmeyen veya yinelenmiş iddia var.")
            if section not in {"short_summary"}:
                values = [item for item in values if item not in used]
                used.update(values)
            result[section] = values
        if not result["short_summary"] or not result["main_idea"]:
            raise RuntimeError("Belge özet finalinin zorunlu bölümleri boş.")
        return _refine_sections(result, claims, profile)
    except RuntimeError:
        return _refine_sections(_fallback_sections(claims, profile), claims, profile)


def _location_label(location: dict[str, Any]) -> str:
    parts: list[str] = []
    if isinstance(location.get("page"), int):
        parts.append(f"s. {location['page']}")
    if location.get("part"):
        labels = {"body": "gövde", "header": "üstbilgi", "footer": "altbilgi"}
        parts.append(labels.get(str(location["part"]), str(location["part"])))
    if isinstance(location.get("paragraph"), int):
        parts.append(f"paragraf {location['paragraph']}")
    if isinstance(location.get("table"), int):
        parts.append(f"tablo {location['table']}")
    if isinstance(location.get("row"), int):
        parts.append(f"satır {location['row']}")
    if isinstance(location.get("column"), int):
        parts.append(f"sütun {location['column']}")
    if isinstance(location.get("line_start"), int):
        end = location.get("line_end")
        parts.append(
            f"satır {location['line_start']}–{end}"
            if isinstance(end, int) and end != location["line_start"]
            else f"satır {location['line_start']}"
        )
    return ", ".join(parts) or "konum kayıtlı"


def _render(
    sections: dict[str, list[str]],
    *,
    claims: list[DocumentClaim],
    blocks: list[dict[str, Any]],
    profile: str,
    source_mode: str,
    source_title: str,
    source_reference: str,
    model: str,
    canonical_sha256: str,
    translation_sha256: str | None,
) -> tuple[str, dict[str, object], dict[str, object]]:
    claim_map = {item.id: item for item in claims}
    block_map = {str(item["id"]): item for item in blocks}
    traces: list[dict[str, object]] = []

    def statement(claim_id: str, section: str) -> str:
        claim = claim_map[claim_id]
        citations = [
            f"{block_id} · {_location_label(block_map[block_id]['location'])}"
            for block_id in claim.block_ids
        ]
        traces.append(
            {
                "statement_id": f"T{len(traces) + 1:04d}",
                "section": section,
                "text": claim.text,
                "claim_ids": [claim.id],
                "source_block_ids": list(claim.block_ids),
                "source_locations": [block_map[item]["location"] for item in claim.block_ids],
                "evidence": [item.to_dict() for item in claim.evidence],
            }
        )
        return f"{claim.text} — " + " ".join(f"[{item}]" for item in citations)

    title_map = {
        "short_summary": "Kısa özet",
        "main_idea": "Ana fikir",
        "important_points": "Önemli noktalar",
        "conclusions": "Sonuçlar",
        "actions": "Yapılacaklar",
    }
    lines = ["# Türkçe Belge Özeti", ""]
    for section in SECTIONS:
        selected = sections.get(section, [])
        if not selected:
            continue
        lines.extend((f"## {title_map[section]}", ""))
        for claim_id in selected:
            lines.append(f"- {statement(claim_id, section)}")
        lines.append("")
    lines.extend(
        (
            "## Kaynak ve üretim bilgisi",
            "",
            f"- Belge: {re.sub(r'[\r\n]+', ' ', source_title).strip()[:500] or 'Bilinmiyor'}",
            f"- Kaynak: {re.sub(r'[\r\n]+', ' ', source_reference).strip()[:500] or 'Yerel belge'}",
            f"- Özet profili: {PROFILE_SETTINGS[profile]['label']}",
            f"- Metin kaynağı: {'doğrulanmış Türkçe çeviri' if source_mode == 'translation' else 'kaynak belge'}",
            f"- Model: `{model}`",
            f"- Kanonik belge SHA-256: `{canonical_sha256}`",
            *(
                (f"- Türkçe çeviri SHA-256: `{translation_sha256}`",)
                if translation_sha256
                else ()
            ),
            "",
        )
    )
    valid_ids = set(block_map)
    supported = sum(
        bool(item["source_block_ids"])
        and all(source_id in valid_ids for source_id in item["source_block_ids"])
        for item in traces
    )
    evidence_valid = sum(
        all(
            evidence["block_id"] in valid_ids
            and _normalized(str(evidence["quote"]))
            in _normalized(str(block_map[str(evidence["block_id"])]["text"]))
            for evidence in item["evidence"]
        )
        for item in traces
    )
    quality = {
        "schema_version": SCHEMA_VERSION,
        "passed": bool(traces) and supported == len(traces) and evidence_valid == len(traces),
        "profile": profile,
        "source_mode": source_mode,
        "statement_count": len(traces),
        "claim_count": len(claims),
        "supported_statement_count": supported,
        "evidence_verified_statement_count": evidence_valid,
        "source_coverage_percent": round(100 * supported / max(len(traces), 1), 2),
        "evidence_coverage_percent": round(100 * evidence_valid / max(len(traces), 1), 2),
        "location_coverage_percent": round(
            100
            * sum(bool(item["source_locations"]) for item in traces)
            / max(len(traces), 1),
            2,
        ),
        "semantic_review_required": True,
        "human_review_status": "pending",
        "errors": [],
        "warnings": [],
    }
    trace = {
        "schema_version": SCHEMA_VERSION,
        "profile": profile,
        "source_mode": source_mode,
        "canonical_jsonl_sha256": canonical_sha256,
        "translation_jsonl_sha256": translation_sha256,
        "source_block_count": len(blocks),
        "claim_count": len(claims),
        "statement_count": len(traces),
        "claims": [item.to_dict() for item in claims],
        "statements": traces,
    }
    return "\n".join(lines), trace, quality


def _write_docx(path: Path, markdown: str) -> None:
    document = Document()
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        elif line.strip():
            document.add_paragraph(line)
    buffer = BytesIO()
    document.save(buffer)
    data = buffer.getvalue()
    Document(BytesIO(data))
    atomic_write_bytes(path, data)


def _write_pdf(path: Path, markdown: str) -> None:
    from PySide6.QtCore import QMarginsF
    from PySide6.QtGui import QGuiApplication, QPageLayout, QPageSize, QTextDocument
    from PySide6.QtPrintSupport import QPrinter

    app = QGuiApplication.instance()
    created = app is None
    if created:
        os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        app = QGuiApplication(["ksi_local-document-summary"])
    temporary = path.with_name(f".{path.name}.partial")
    temporary.unlink(missing_ok=True)
    try:
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(temporary))
        printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        printer.setPageMargins(QMarginsF(18, 18, 18, 18), QPageLayout.Unit.Millimeter)
        document = QTextDocument()
        document.setMarkdown(markdown)
        document.print_(printer)
        if not temporary.is_file() or not temporary.read_bytes().startswith(b"%PDF"):
            raise RuntimeError("Belge özeti PDF olarak üretilemedi.")
        PdfReader(temporary)
        os.replace(temporary, path)
    except Exception:
        temporary.unlink(missing_ok=True)
        raise
    finally:
        if created:
            app.quit()


def _checkpoint_records(
    path: Path,
    *,
    canonical_sha256: str,
    translation_sha256: str | None,
    source_mode: str,
    profile: str,
    model: str,
    chunks: list[list[dict[str, Any]]],
) -> tuple[int, list[DocumentClaim]]:
    if not path.is_file():
        return 0, []
    if path.is_symlink():
        raise ValueError("Belge özet checkpoint'i sembolik bağlantı olamaz.")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError("Belge özet checkpoint'i okunamadı.") from error
    if not isinstance(payload, dict) or payload.get("schema_version") != SCHEMA_VERSION:
        raise ValueError("Belge özet checkpoint'i kaynak, profil veya modelle eşleşmiyor.")
    if payload.get("summary_policy_version") != SUMMARY_POLICY_VERSION:
        return 0, []
    if (
        payload.get("canonical_jsonl_sha256") != canonical_sha256
        or payload.get("source_mode") != source_mode
        or payload.get("profile") != profile
        or payload.get("model") != model
    ):
        raise ValueError("Belge özet checkpoint'i kaynak, profil veya modelle eşleşmiyor.")
    # A revised, still-valid translation is a normal upstream change. Its summary
    # checkpoint is stale rather than corrupt, so regenerate it from the beginning.
    if payload.get("translation_jsonl_sha256") != translation_sha256:
        return 0, []
    completed = payload.get("completed_chunks")
    records = payload.get("claims")
    if not isinstance(completed, int) or not 0 <= completed <= len(chunks) or not isinstance(records, list):
        raise ValueError("Belge özet checkpoint parça sayısı veya iddiaları geçersiz.")
    allowed = {str(item["id"]) for chunk in chunks[:completed] for item in chunk}
    completed_text = {
        str(item["id"]): str(item["text"])
        for chunk in chunks[:completed]
        for item in chunk
    }
    claims: list[DocumentClaim] = []
    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict) or record.get("id") != f"C{index:06d}":
            raise ValueError("Belge özet checkpoint iddia sırası geçersiz.")
        try:
            evidence = tuple(
                DocumentEvidence(str(item["block_id"]), str(item["quote"]))
                for item in record["evidence"]
            )
            grounded_text, grounded_category = _ground_claim_text(
                str(record["text"]), str(record["category"]), evidence
            )
            claim = DocumentClaim(
                str(record["id"]),
                grounded_text,
                grounded_category,
                tuple(str(item) for item in record["block_ids"]),
                evidence,
            )
        except (KeyError, TypeError) as error:
            raise ValueError("Belge özet checkpoint iddiası okunamadı.") from error
        if (
            claim.category not in CATEGORIES
            or not claim.text.strip()
            or not claim.block_ids
            or any(item not in allowed for item in claim.block_ids)
            or {item.block_id for item in claim.evidence} != set(claim.block_ids)
            or any(
                not item.quote.strip()
                or _normalized(item.quote)
                not in _normalized(completed_text.get(item.block_id, ""))
                for item in claim.evidence
            )
        ):
            raise ValueError("Belge özet checkpoint kaynak kimlikleri geçersiz.")
        claims.append(claim)
    return completed, claims


def summarize_document(
    canonical_path: str | Path,
    output_directory: str | Path,
    *,
    client: OllamaClient,
    source_mode: str = "auto",
    profile: str = "standard",
    translation_path: str | Path | None = None,
    translation_quality_path: str | Path | None = None,
    source_title: str = "Belge",
    source_reference: str = "Yerel belge",
    model: str = SUMMARY_MODEL,
    max_chars: int = 8000,
    checkpoint_path: str | Path | None = None,
    create_pdf: bool = False,
    on_progress: Callable[[DocumentSummaryProgress], None] | None = None,
    unload_model_at_end: bool = True,
) -> DocumentSummaryResult:
    if profile not in PROFILE_SETTINGS:
        raise ValueError("Belge özet profili kısa, standart veya ayrıntılı olmalıdır.")
    resolved_mode, canonical_sha256, translation_sha256, blocks = _summary_blocks(
        canonical_path,
        source_mode=source_mode,
        translation_path=translation_path,
        translation_quality_path=translation_quality_path,
    )
    output = Path(output_directory).expanduser()
    if output.is_symlink():
        raise ValueError("Belge özet çıktı klasörü sembolik bağlantı olamaz.")
    output.mkdir(parents=True, exist_ok=True, mode=0o700)
    output = output.resolve()
    checkpoint = (
        Path(checkpoint_path).expanduser()
        if checkpoint_path
        else output / ".belge-ozeti.checkpoint.json"
    )
    if checkpoint.is_symlink():
        raise ValueError("Belge özet checkpoint'i sembolik bağlantı olamaz.")
    checkpoint.parent.mkdir(parents=True, exist_ok=True, mode=0o700)
    chunks = split_document_blocks(
        blocks,
        max_chars=max_chars,
        max_blocks=int(PROFILE_SETTINGS[profile]["claims_per_chunk"]),
    )
    completed, claims = _checkpoint_records(
        checkpoint,
        canonical_sha256=canonical_sha256,
        translation_sha256=translation_sha256,
        source_mode=resolved_mode,
        profile=profile,
        model=model,
        chunks=chunks,
    )
    total_steps = len(chunks) + 1
    for index, chunk in enumerate(chunks[completed:], start=completed):
        new_claims = _generate_claims(
            chunk,
            client=client,
            model=model,
            profile=profile,
            chunk_index=index,
            chunk_count=len(chunks),
            first_number=len(claims) + 1,
        )
        claims.extend(new_claims)
        atomic_write_json(
            checkpoint,
            {
                "schema_version": SCHEMA_VERSION,
                "summary_policy_version": SUMMARY_POLICY_VERSION,
                "canonical_jsonl_sha256": canonical_sha256,
                "translation_jsonl_sha256": translation_sha256,
                "source_mode": resolved_mode,
                "profile": profile,
                "model": model,
                "completed_chunks": index + 1,
                "complete": False,
                "claims": [item.to_dict() for item in claims],
            },
        )
        if on_progress:
            on_progress(DocumentSummaryProgress(index + 1, total_steps))
    sections = _select_sections(
        claims,
        client=client,
        model=model,
        profile=profile,
        unload_model_at_end=unload_model_at_end,
    )
    markdown, trace, quality = _render(
        sections,
        claims=claims,
        blocks=blocks,
        profile=profile,
        source_mode=resolved_mode,
        source_title=source_title,
        source_reference=source_reference,
        model=model,
        canonical_sha256=canonical_sha256,
        translation_sha256=translation_sha256,
    )
    markdown_path = output / "belge-ozeti.md"
    docx_path = output / "belge-ozeti.docx"
    pdf_path = output / "belge-ozeti.pdf" if create_pdf else None
    trace_path = output / "belge-ozeti.kaynaklar.json"
    quality_path = output / "belge-ozeti.kalite.json"
    staging = Path(tempfile.mkdtemp(prefix=".belge-ozeti-", dir=output))
    try:
        staged_markdown = staging / markdown_path.name
        staged_docx = staging / docx_path.name
        staged_trace = staging / trace_path.name
        staged_quality = staging / quality_path.name
        atomic_write_text(staged_markdown, markdown.rstrip() + "\n")
        _write_docx(staged_docx, markdown)
        staged_files = [
            (staged_markdown, markdown_path),
            (staged_docx, docx_path),
        ]
        if pdf_path is not None:
            staged_pdf = staging / pdf_path.name
            _write_pdf(staged_pdf, markdown)
            staged_files.append((staged_pdf, pdf_path))
        atomic_write_json(staged_trace, trace)
        atomic_write_json(staged_quality, quality)
        staged_files.extend(((staged_trace, trace_path), (staged_quality, quality_path)))
        for source, target in staged_files:
            os.replace(source, target)
    finally:
        shutil.rmtree(staging, ignore_errors=True)
    atomic_write_json(
        checkpoint,
        {
            "schema_version": SCHEMA_VERSION,
            "summary_policy_version": SUMMARY_POLICY_VERSION,
            "canonical_jsonl_sha256": canonical_sha256,
            "translation_jsonl_sha256": translation_sha256,
            "source_mode": resolved_mode,
            "profile": profile,
            "model": model,
            "completed_chunks": len(chunks),
            "complete": True,
            "claims": [item.to_dict() for item in claims],
        },
    )
    if on_progress:
        on_progress(DocumentSummaryProgress(total_steps, total_steps))
    return DocumentSummaryResult(
        markdown_path=markdown_path,
        docx_path=docx_path,
        pdf_path=pdf_path,
        trace_path=trace_path,
        quality_path=quality_path,
        checkpoint_path=checkpoint,
        profile=profile,
        source_mode=resolved_mode,
        chunk_count=len(chunks),
        claim_count=len(claims),
        statement_count=int(quality["statement_count"]),
        quality=quality,
    )
