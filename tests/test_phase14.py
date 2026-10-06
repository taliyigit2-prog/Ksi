from __future__ import annotations

import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from PySide6.QtWidgets import QApplication
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from ksi_local.document_extraction import (
    LOW_OCR_CONFIDENCE,
    extract_document,
    vision_capabilities,
)
from ksi_local.gui import MainWindow
from ksi_local.job_store import JobKind, JobStatus
from ksi_local.settings import WorkspacePaths
from ksi_local.system_health import build_health_report


def workspace_at(root: Path) -> WorkspacePaths:
    workspace = root / "KSI-Workspace"
    paths = WorkspacePaths(
        root=workspace,
        jobs=workspace / "jobs",
        outputs=workspace / "outputs",
        models_ollama=workspace / "models/ollama",
        models_whisper=workspace / "models/whisper",
        yt_dlp=workspace / "tools/yt-dlp/unused/yt-dlp",
        deno=workspace / "tools/deno/unused/deno",
    )
    paths.jobs.mkdir(parents=True)
    return paths


def text_pdf(path: Path, pages: list[list[str]]) -> Path:
    writer = PdfWriter()
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_reference = writer._add_object(font)
    for lines in pages:
        page = writer.add_blank_page(width=612, height=792)
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_reference})}
        )
        commands = ["BT /F1 12 Tf 72 740 Td"]
        for index, line in enumerate(lines):
            if index:
                commands.append("0 -36 Td")
            escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({escaped}) Tj")
        commands.append("ET")
        stream = DecodedStreamObject()
        stream.set_data(" ".join(commands).encode("latin-1"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)
    return path


class CanonicalExtractionTests(unittest.TestCase):
    def test_single_paragraph_txt_with_terminal_newline_is_extracted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "single.txt"
            source.write_text(
                "El documento termina con una nueva línea y permanece íntegro.\n",
                encoding="utf-8",
            )
            result = extract_document(source, root / "work")
            records = [
                json.loads(line)
                for line in result.jsonl_path.read_text(encoding="utf-8").splitlines()
            ]
            self.assertEqual(result.block_count, 1)
            self.assertEqual(
                records[0]["text"],
                "El documento termina con una nueva línea y permanece íntegro.",
            )

    def test_txt_writes_traceable_jsonl_text_and_quality(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "notes.txt"
            source.write_text(
                "This is the first sufficiently long English paragraph for detection.\n\n"
                "This is the second paragraph and it must remain in source order.",
                encoding="utf-8",
            )
            result = extract_document(source, root / "work")
            records = [
                json.loads(line)
                for line in result.jsonl_path.read_text(encoding="utf-8").splitlines()
            ]
            self.assertEqual([item["id"] for item in records], ["B000001", "B000002"])
            self.assertEqual([item["sequence"] for item in records], [1, 2])
            self.assertTrue(all(item["source_sha256"] == result.source_sha256 for item in records))
            self.assertEqual(records[0]["location"]["line_start"], 1)
            self.assertEqual(records[1]["location"]["line_start"], 3)
            self.assertEqual(result.quality["document_language"], "en")
            self.assertTrue(result.quality["passed"])
            self.assertIn("[B000001]", result.text_path.read_text(encoding="utf-8"))

    def test_markdown_separates_heading_list_code_table_and_link_targets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "guide.md"
            source.write_text(
                "# Local Guide\n\n"
                "Read the [documentation](https://example.invalid/docs).\n\n"
                "- First safe item\n"
                "- Second safe item\n\n"
                "```python\nprint('never executed')\n```\n\n"
                "| Name | State |\n| --- | --- |\n| OCR | Local |\n",
                encoding="utf-8",
            )
            result = extract_document(source, root / "work")
            records = [json.loads(line) for line in result.jsonl_path.read_text().splitlines()]
            types = [item["block_type"] for item in records]
            self.assertIn("heading", types)
            self.assertIn("list_item", types)
            self.assertIn("code_block", types)
            self.assertIn("table_cell", types)
            paragraph = next(item for item in records if item["block_type"] == "paragraph")
            self.assertEqual(
                paragraph["style"]["links"][0]["target"],
                "https://example.invalid/docs",
            )
            code = next(item for item in records if item["block_type"] == "code_block")
            self.assertEqual(code["text"], "print('never executed')")

    def test_non_utf8_txt_is_rejected_with_actionable_message(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "legacy.txt"
            source.write_bytes("café français".encode("cp1252"))
            with self.assertRaisesRegex(ValueError, "UTF-8"):
                extract_document(source, Path(directory) / "work")

    def test_source_output_and_ocr_helper_symlinks_are_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "real.txt"
            source.write_text("A sufficiently long local document source for testing.")
            source_link = root / "source.txt"
            source_link.symlink_to(source)
            with self.assertRaisesRegex(ValueError, "sembolik"):
                extract_document(source_link, root / "work")

            output = root / "actual-output"
            output.mkdir()
            output_link = root / "linked-output"
            output_link.symlink_to(output, target_is_directory=True)
            with self.assertRaisesRegex(ValueError, "sembolik"):
                extract_document(source, output_link)

            helper_link = root / "linked-helper"
            helper_link.symlink_to(Path(__file__).parents[1] / "build/KSIOCR")
            with self.assertRaisesRegex(RuntimeError, "sembolik"):
                vision_capabilities(helper_link)


class PDFExtractionTests(unittest.TestCase):
    def test_native_pdf_preserves_pages_and_flags_repeated_margins(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = text_pdf(
                root / "native.pdf",
                [
                    [
                        "Repeated Header",
                        f"Page {index} has enough English body text for extraction.",
                        "Repeated Footer",
                    ]
                    for index in range(1, 4)
                ],
            )
            result = extract_document(source, root / "work")
            records = [json.loads(line) for line in result.jsonl_path.read_text().splitlines()]
            self.assertEqual(result.quality["details"]["native_text_pages"], [1, 2, 3])
            self.assertEqual(result.quality["details"]["ocr_pages"], [])
            self.assertTrue(any("repeated_header" in item["flags"] for item in records))
            self.assertTrue(any("repeated_footer" in item["flags"] for item in records))
            self.assertEqual({item["location"]["page"] for item in records}, {1, 2, 3})

    def test_scanned_pdf_uses_ocr_confidence_and_bounding_box(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with source.open("wb") as handle:
                writer.write(handle)
            mocked_lines = [
                {
                    "text": "Recognized local OCR line",
                    "confidence": LOW_OCR_CONFIDENCE - 0.1,
                    "x": 0.1,
                    "y": 0.8,
                    "width": 0.7,
                    "height": 0.05,
                }
            ]
            capabilities = {
                "engine": "macOS Vision",
                "revision": 3,
                "supported_languages": ("en-US", "tr-TR"),
                "automatic_language_detection": True,
            }
            with patch(
                "ksi_local.document_extraction._ocr_pdf_page",
                return_value=mocked_lines,
            ), patch(
                "ksi_local.document_extraction.vision_capabilities",
                return_value=capabilities,
            ):
                result = extract_document(source, root / "work", ocr_helper=root / "fake")
            record = json.loads(result.jsonl_path.read_text().splitlines()[0])
            self.assertEqual(record["extraction_method"], "vision_ocr")
            self.assertIn("low_ocr_confidence", record["flags"])
            self.assertEqual(record["location"]["bounding_box"]["y"], 0.8)
            self.assertEqual(result.quality["low_ocr_confidence_block_count"], 1)


class DOCXExtractionTests(unittest.TestCase):
    def test_distinct_table_cells_are_not_lost_when_wrappers_are_reused(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "table.docx"
            document = Document()
            table = document.add_table(rows=3, cols=3)
            expected = []
            for row_index, row in enumerate(table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    value = f"R{row_index}C{column_index} value"
                    cell.text = value
                    expected.append(value)
            document.save(source)
            result = extract_document(source, root / "work")
            records = [json.loads(line) for line in result.jsonl_path.read_text().splitlines()]
            cells = [item["text"] for item in records if item["block_type"] == "table_cell"]
            self.assertEqual(cells, expected)

    def test_docx_body_table_header_footer_and_footnote_are_located(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "structured.docx"
            document = Document()
            document.add_heading("Document heading", level=1)
            document.add_paragraph(
                "This English body paragraph is long enough for offline language detection."
            )
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Name"
            table.cell(0, 1).text = "Value"
            table.cell(1, 0).text = "Mode"
            table.cell(1, 1).text = "Local"
            document.sections[0].header.paragraphs[0].text = "Header text"
            document.sections[0].footer.paragraphs[0].text = "Footer text"
            document.save(source)
            footnotes = (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:footnote w:id="1"><w:p><w:r><w:t>Footnote text</w:t></w:r></w:p></w:footnote>'
                '</w:footnotes>'
            )
            with zipfile.ZipFile(source, "a", zipfile.ZIP_DEFLATED) as package:
                package.writestr("word/footnotes.xml", footnotes)
            result = extract_document(source, root / "work")
            records = [json.loads(line) for line in result.jsonl_path.read_text().splitlines()]
            self.assertEqual(records[0]["block_type"], "heading")
            self.assertTrue(any(item["block_type"] == "table_cell" for item in records))
            self.assertTrue(any(item["location"].get("part") == "header" for item in records))
            self.assertTrue(any(item["location"].get("part") == "footer" for item in records))
            self.assertTrue(any(item["block_type"] == "footnote" for item in records))
            self.assertEqual(result.quality["details"]["footnote_block_count"], 1)


class VisionAndQueueTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.application = QApplication.instance() or QApplication([])

    def test_compiled_vision_helper_reports_required_languages(self) -> None:
        helper = Path(__file__).resolve().parents[1] / "build/KSIOCR"
        if not helper.is_file():
            self.skipTest("Vision helper has not been built")
        result = vision_capabilities(helper)
        languages = set(result["supported_languages"])
        self.assertTrue({"en-US", "ru-RU", "de-DE", "fr-FR", "it-IT"} <= languages)
        self.assertTrue({"es-ES", "zh-Hans", "zh-Hant"} & languages)
        self.assertTrue(result["automatic_language_detection"])

    def test_system_health_includes_local_vision_helper(self) -> None:
        helper = Path(__file__).resolve().parents[1] / "build/KSIOCR"
        if not helper.is_file():
            self.skipTest("Vision helper has not been built")
        status = next(
            item for item in build_health_report().tools if item.name == "Apple Vision OCR"
        )
        self.assertTrue(status.installed)
        self.assertIn("Vision revizyon", status.version or "")
        self.assertIn("dil varyantı", status.note or "")

    def test_imported_document_is_queued_for_extraction_without_model(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workspace = workspace_at(root)
            job = workspace.jobs / "JOB-PHASE14"
            source = job / "source/original.txt"
            source.parent.mkdir(parents=True)
            source.write_text("Local document source", encoding="utf-8")
            import hashlib

            digest = hashlib.sha256(source.read_bytes()).hexdigest()
            (job / "manifest.json").write_text(
                json.dumps({"import": {"sha256": digest}}), encoding="utf-8"
            )
            with (
                patch.dict(os.environ, {"KSI_STATE_DIRECTORY": str(root / "state")}),
                patch("ksi_local.gui.resolve_workspace", return_value=workspace),
            ):
                window = MainWindow()
                window.workspace = workspace
                record = window.store.create_job(
                    job_id="JOB-PHASE14",
                    job_kind=JobKind.DOCUMENT,
                    source=str(source),
                    source_language="auto",
                    want_subtitle=True,
                    want_summary=True,
                    want_dub=False,
                    job_directory=job,
                )
                window.store.transition_job(record.id, JobStatus.RUNNING)
                window.store.transition_job(record.id, JobStatus.IMPORTED)
                record = window.store.get_job(record.id)
                with patch.object(window, "_run_next"):
                    window._launch_document_extraction(record)
                self.assertEqual(window.store.get_job(record.id).status, JobStatus.QUEUED)
                self.assertEqual(window.pending[0].stage, "document_extract")
                command = window.pending[0].argv
                self.assertIn("ksi_local.document_worker", command)
                self.assertNotIn("ollama", " ".join(command).casefold())
                window.close()


if __name__ == "__main__":
    unittest.main()
