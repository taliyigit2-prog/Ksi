from __future__ import annotations

import json
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from docx import Document
from pypdf import PdfReader
from PySide6.QtWidgets import QApplication

from ksi_local.document_output import create_document_outputs
from ksi_local.document_translation import translate_document
from ksi_local.exporter import export_artifacts, select_job_artifacts, sha256_file
from ksi_local.gui import MainWindow
from ksi_local.job_store import JobKind, JobStatus
from ksi_local.settings import WorkspacePaths


TEST_APP = QApplication.instance() or QApplication([])


def records() -> list[dict[str, object]]:
    return [
        {
            "id": "B000001",
            "block_type": "heading",
            "translated_text": "Türkçe Başlık: ığüşöçİ",
            "location": {"part": "body", "paragraph": 1},
            "style": {"heading_level": 1},
        },
        {
            "id": "B000002",
            "block_type": "paragraph",
            "translated_text": "Bu belge, Türkçe karakterleri ve anlamlı içeriği eksiksiz korur.",
            "location": {"part": "body", "paragraph": 2},
            "style": {},
        },
        {
            "id": "B000003",
            "block_type": "list_item",
            "translated_text": "Birinci liste maddesi",
            "location": {"part": "body", "paragraph": 3},
            "style": {"indent": 0},
        },
        {
            "id": "B000004",
            "block_type": "table_cell",
            "translated_text": "Sütun",
            "location": {"part": "body", "table": 1, "row": 1, "column": 1},
            "style": {"header": True},
        },
        {
            "id": "B000005",
            "block_type": "table_cell",
            "translated_text": "Değer",
            "location": {"part": "body", "table": 1, "row": 2, "column": 1},
            "style": {},
        },
        {
            "id": "B000006",
            "block_type": "code_block",
            "translated_text": "print('kaynak kod korunur')",
            "location": {"part": "body", "paragraph": 4},
            "style": {"language": "python"},
        },
    ]


def write_turkish_canonical(path: Path) -> Path:
    payloads = []
    for sequence, record in enumerate(records(), start=1):
        text = str(record["translated_text"])
        payloads.append(
            {
                "schema_version": 1,
                "id": f"B{sequence:06d}",
                "sequence": sequence,
                "source_sha256": "a" * 64,
                "source_format": "docx",
                "block_type": record["block_type"],
                "text": text,
                "location": record["location"],
                "extraction_method": "test",
                "detected_language": "tr",
                "language_confidence": 1.0,
                "ocr_confidence": None,
                "style": record["style"],
                "flags": [],
            }
        )
    path.write_text(
        "\n".join(json.dumps(item, ensure_ascii=False) for item in payloads) + "\n",
        encoding="utf-8",
    )
    return path


def workspace_at(root: Path) -> WorkspacePaths:
    workspace = WorkspacePaths(
        root=root / "KSI-Workspace",
        jobs=root / "KSI-Workspace/jobs",
        outputs=root / "KSI-Workspace/outputs",
        models_ollama=root / "KSI-Workspace/models/ollama",
        models_whisper=root / "KSI-Workspace/models/whisper",
        yt_dlp=root / "KSI-Workspace/tools/yt-dlp/yt-dlp",
        deno=root / "KSI-Workspace/tools/deno/deno",
    )
    workspace.jobs.mkdir(parents=True)
    return workspace


class DocumentOutputTests(unittest.TestCase):
    def test_markdown_docx_pdf_and_txt_are_readable_and_complete(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            result = create_document_outputs(records(), directory, source_title="örnek.docx")
            markdown = result.markdown_path.read_text(encoding="utf-8")
            plain = result.text_path.read_text(encoding="utf-8")
            for record in records():
                value = str(record["translated_text"])
                self.assertIn(value, markdown)
                self.assertIn(value, plain)
            self.assertIn("# Türkçe Başlık", markdown)
            self.assertIn("- Birinci liste maddesi", markdown)
            self.assertIn("| Sütun |", markdown)
            self.assertIn("```python", markdown)
            document = Document(result.docx_path)
            docx_text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [
                    cell.text
                    for table in document.tables
                    for row in table.rows
                    for cell in row.cells
                ]
            )
            for value in ("Türkçe Başlık", "Birinci liste maddesi", "Sütun", "Değer"):
                self.assertIn(value, docx_text)
            self.assertGreater(len(PdfReader(result.pdf_path, strict=True).pages), 0)
            self.assertTrue(result.embedded_pdf_font)
            for path in (result.markdown_path, result.docx_path, result.pdf_path, result.text_path):
                self.assertEqual(path.stat().st_mode & 0o777, 0o600)

    def test_output_failure_leaves_no_visible_partial_artifact(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            with patch("ksi_local.document_output._write_pdf", side_effect=OSError("disk lost")):
                with self.assertRaises(OSError):
                    create_document_outputs(records(), root, source_title="örnek")
            self.assertEqual(list(root.glob("belge-turkce.*")), [])
            self.assertEqual(list(root.glob(".belge-turkce-*")), [])

    def test_translation_creates_four_formats_and_hashes(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            canonical = write_turkish_canonical(root / "source.jsonl")
            result = translate_document(
                canonical,
                root / "outputs",
                client=None,
                source_title="Türkçe örnek.docx",
            )
            self.assertTrue(result.quality["passed"])
            self.assertTrue(result.quality["pdf_font_embedded"])
            for key, path in (
                ("text_sha256", result.text_path),
                ("markdown_sha256", result.markdown_path),
                ("docx_sha256", result.docx_path),
                ("pdf_sha256", result.pdf_path),
            ):
                self.assertEqual(result.quality[key], sha256_file(path))


class DocumentExportTests(unittest.TestCase):
    def test_two_gib_reserve_and_symlink_target_are_enforced(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "belge.pdf"
            source.write_bytes(b"pdf")
            target = root / "target"
            target.mkdir()
            usage_type = type(shutil.disk_usage(target))
            with patch(
                "ksi_local.exporter.shutil.disk_usage",
                return_value=usage_type(total=10, used=9, free=1),
            ):
                with self.assertRaisesRegex(OSError, "2 GiB"):
                    export_artifacts([source], desktop=target, folder_name="Belge")
            link = root / "target-link"
            link.symlink_to(target, target_is_directory=True)
            with self.assertRaisesRegex(ValueError, "sembolik"):
                export_artifacts([source], desktop=link, folder_name="Belge")

    def test_document_results_copy_with_manifest_collision_and_source_preserved(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            job = root / "job"
            outputs = job / "outputs"
            source = job / "source"
            outputs.mkdir(parents=True)
            source.mkdir()
            original = source / "original.docx"
            original.write_bytes(b"original")
            for suffix in (".txt", ".md", ".docx", ".pdf"):
                (outputs / f"belge-turkce{suffix}").write_bytes(f"translation{suffix}".encode())
            for suffix in (".md", ".docx", ".pdf"):
                (outputs / f"belge-ozeti{suffix}").write_bytes(f"summary{suffix}".encode())
            target = root / "Documents"
            target.mkdir()
            artifacts = select_job_artifacts(job, export_kind="all", title="Örnek Belge")
            first = export_artifacts(artifacts, desktop=target, folder_name="Örnek Belge")
            second = export_artifacts(artifacts, desktop=target, folder_name="Örnek Belge")
            self.assertEqual(first.name, "Örnek Belge")
            self.assertEqual(second.name, "Örnek Belge (2)")
            manifest = json.loads((first / "kopya-manifest.json").read_text())
            self.assertEqual(len(manifest["files"]), 7)
            for item in manifest["files"]:
                self.assertEqual(item["source_sha256"], item["copied_sha256"])
                self.assertEqual(item["copied_sha256"], sha256_file(first / item["filename"]))
            self.assertEqual(original.read_bytes(), b"original")

    def test_gui_exposes_three_mac_targets_and_document_job_can_complete(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            workspace = workspace_at(root)
            job = workspace.jobs / "JOB-PHASE17"
            (job / "outputs").mkdir(parents=True)
            (job / "outputs/belge-turkce.jsonl").write_text("{}\n")
            for suffix in (".txt", ".md", ".docx", ".pdf"):
                (job / f"outputs/belge-turkce{suffix}").write_bytes(b"ready")
            (job / "outputs/belge-turkce.kalite.json").write_text("{}")
            with (
                patch.dict(os.environ, {"KSI_STATE_DIRECTORY": str(root / "state")}),
                patch("ksi_local.gui.resolve_workspace", return_value=workspace),
            ):
                window = MainWindow()
                window.workspace = workspace
                record = window.store.create_job(
                    job_id="JOB-PHASE17",
                    job_kind=JobKind.DOCUMENT,
                    source=str(job / "source/original.docx"),
                    source_language="tr",
                    want_subtitle=True,
                    want_summary=False,
                    want_dub=False,
                    job_directory=job,
                )
                window.store.transition_job(record.id, JobStatus.RUNNING)
                window.current_job_id = record.id
                window.current_job = job
                window._completed()
                self.assertEqual(window.store.get_job(record.id).status, JobStatus.COMPLETED)
                self.assertEqual(
                    [window.export_target.itemData(index) for index in range(3)],
                    ["desktop", "documents", "custom"],
                )
                self.assertEqual(window.export_button.text(), "Mac'e Kopyala")
                window.close()


if __name__ == "__main__":
    unittest.main()
